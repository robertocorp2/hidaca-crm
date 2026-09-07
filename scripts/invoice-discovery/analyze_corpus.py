#!/usr/bin/env python3
"""Read-only HIDACA invoice corpus profiler.

The script never writes to the source tree. It hashes every file and writes
machine-readable discovery outputs to a caller-provided analysis directory.
"""

from __future__ import annotations

import argparse
import concurrent.futures
import hashlib
import json
import re
import statistics
import sys
import traceback
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader


PARSER_VERSION = "hidaca-invoice-discovery/1.0.0"
INVOICE_EXTENSIONS = {".xlsx", ".pdf"}
RESTRICTED_EXTENSIONS = {".p12"}
UNSUPPORTED_EXTENSIONS = {".lnk"}

FIELD_PATTERNS = [
    ("ncf", re.compile(r"\bNCF\s*:\s*(.+)", re.I)),
    ("issue_date", re.compile(r"\bFecha\s*:\s*(.+)", re.I)),
    ("customer_name", re.compile(r"\bCliente\s*:\s*(.+)", re.I)),
    ("customer_tax_id", re.compile(r"\bRNC(?:/C[eé]dula)?\s*:\s*(.+)", re.I)),
    ("invoice_number", re.compile(r"\bFactura\s*(?:No\.?|N[úu]m(?:ero)?)?\s*:\s*(.+)", re.I)),
    ("customer_address", re.compile(r"\bDirecci[oó]n\s*:\s*(.+)", re.I)),
    ("project_name", re.compile(r"\bProyecto\s*:\s*(.+)", re.I)),
    ("phone", re.compile(r"\b(?:Tel[eé]fonos?|Tel\.?|Celular|M[oó]vil)\s*:\s*(.+)", re.I)),
    ("contact_name", re.compile(r"\bContacto\s*:\s*(.+)", re.I)),
    ("email", re.compile(r"\b(?:Correo|E-?mail)\s*:\s*(.+)", re.I)),
    ("due_date", re.compile(r"\bVencimiento\s*:\s*(.+)", re.I)),
    ("payment_terms", re.compile(r"\bCondiciones?\s+de\s+Pago\s*:\s*(.+)", re.I)),
    ("quotation_number", re.compile(r"\bCotizaci[oó]n\s*:\s*(.+)", re.I)),
    ("purchase_order", re.compile(r"\bOrden\s+de\s+Compra\s*:\s*(.+)", re.I)),
    ("representative", re.compile(r"\bRepresentante\s*:\s*(.+)", re.I)),
]

STOP_LABELS = [
    "Factura No.",
    "Factura No:",
    "Factura:",
    "Condiciones de Pago:",
    "Condición:",
    "Cotización:",
    "Cotizacion:",
    "Orden de Compra:",
    "Representante:",
    "Vencimiento:",
    "Proyecto:",
    "Cliente:",
    "Dirección:",
    "Direccion:",
    "Contacto:",
    "Correo:",
    "Celular:",
    "Teléfonos:",
    "Telefonos:",
    "Fecha:",
    "RNC:",
    "NCF:",
]

TOTAL_LABELS = {
    "sub-total": "subtotal",
    "subtotal": "subtotal",
    "sub total": "subtotal",
    "itbis": "tax_amount",
    "itbis 18%": "tax_amount",
    "total (rd$)": "total_amount",
    "total": "total_amount",
    "avance": "advance_amount",
    "pendiente": "balance_amount",
    "reparacion": "repair_amount",
    "reparación": "repair_amount",
    "instalacion": "installation_amount",
    "instalación": "installation_amount",
}

STANDALONE_LABELS = {
    "no",
    "fecha",
    "mes",
    "año",
    "ano",
    "cliente",
    "factura",
    "ncf",
    "tipo",
    "sub-total",
    "subtotal",
    "sub total",
    "itbis",
    "itbis 18%",
    "total",
    "total (rd$)",
    "fecha cierre",
    "avance",
    "pendiente",
    "estatus",
    "descripción",
    "descripcion",
    "codigo",
    "código",
    "cant",
    "cantidad",
    "ubicación",
    "ubicacion",
    "ancho (cm)",
    "altura (cm)",
    "área (m2)",
    "area (m2)",
    "precio",
    "reparacion",
    "reparación",
    "instalacion",
    "instalación",
    "shutter",
    "valores",
    "realizado por",
    "autorizado por",
    "condición",
    "condicion",
    "correo",
    "contacto",
    "proyecto",
    "cotización",
    "cotizacion",
    "orden de compra",
    "representante",
    "vencimiento",
}


def json_default(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return str(value)


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=json_default),
        encoding="utf-8",
    )


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def scalar(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, Decimal)):
        return value.isoformat() if hasattr(value, "isoformat") else str(value)
    return str(value)


def normalize_label(value: str) -> str:
    text = re.sub(r"\s+", " ", value.replace("\n", " ")).strip(" :.-")
    return text.casefold()


def split_value_at_next_label(value: str) -> str:
    text = re.sub(r"\s+", " ", value).strip()
    positions = []
    for label in STOP_LABELS:
        match = re.search(re.escape(label), text, re.I)
        if match and match.start() > 0:
            positions.append(match.start())
    if positions:
        text = text[: min(positions)]
    return text.strip(" ;|")


def parse_decimal(value: Any) -> Decimal | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    text = str(value).strip()
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"[^\d,.\-]", "", text).replace(",", "")
    try:
        parsed = Decimal(cleaned)
        return -parsed if negative and parsed > 0 else parsed
    except InvalidOperation:
        return None


def extract_inline_fields(text: str) -> dict[str, list[str]]:
    found: dict[str, list[str]] = defaultdict(list)
    lines = text.splitlines() or [text]
    for line in lines:
        compact = re.sub(r"\s+", " ", line).strip()
        if not compact:
            continue
        for key, pattern in FIELD_PATTERNS:
            for match in pattern.finditer(compact):
                value = split_value_at_next_label(match.group(1))
                if value and value not in found[key]:
                    found[key].append(value)
    return dict(found)


def label_candidates_from_text(text: str) -> list[str]:
    candidates: list[str] = []
    for line in text.splitlines():
        line = re.sub(r"\s+", " ", line).strip()
        if not line:
            continue
        for match in re.finditer(r"(?P<label>[A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9 /().$%-]{2,45})\s*:", line):
            label = match.group("label").strip()
            if 2 <= len(label) <= 45:
                candidates.append(label)
    return candidates


def discover_source_paths(acquired_root: Path, matrix_path: Path, restricted_root: Path) -> list[tuple[Path, str]]:
    discovered: list[tuple[Path, str]] = []
    for path in sorted(p for p in acquired_root.rglob("*") if p.is_file()):
        discovered.append((path, f"FACTURAS HIDACA/{path.relative_to(acquired_root).as_posix()}"))
    discovered.append((matrix_path, "FACTURAS HIDACA/FACTURACION HIDACA - Matriz Principal.xlsx"))
    if restricted_root.exists():
        for path in sorted(p for p in restricted_root.rglob("*") if p.is_file()):
            discovered.append((path, f"FACTURAS HIDACA/FACTURADOR DGII/{path.relative_to(restricted_root).as_posix()}"))
    return discovered


def base_manifest(path: Path, relative_path: str) -> dict[str, Any]:
    stat = path.stat()
    ext = path.suffix.lower()
    if ext in RESTRICTED_EXTENSIONS:
        category = "irrelevant"
        status = "restricted_inventory_only"
        reason = "Sensitive certificate/private-key container; hashed and inventoried only."
    elif ext in UNSUPPORTED_EXTENSIONS:
        category = "unsupported"
        status = "unsupported"
        reason = "Windows shortcut; never executed or dereferenced."
    elif ext in INVOICE_EXTENSIONS:
        category = "business_relevant"
        status = "pending_parse"
        reason = None
    else:
        category = "unsupported"
        status = "unsupported"
        reason = f"Unsupported extension {ext or '[none]'}."
    return {
        "relative_path": relative_path,
        "file_name": path.name,
        "extension": ext,
        "size_bytes": stat.st_size,
        "source_modified_at": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(),
        "download_status": "downloaded",
        "sha256": sha256_file(path),
        "category": category,
        "parse_status": status,
        "reason": reason,
        "local_path": str(path),
    }


def parse_pdf(path: Path, relative_path: str) -> dict[str, Any]:
    profile: dict[str, Any] = {
        "relative_path": relative_path,
        "parser": "pypdf.PdfReader",
        "parser_version": f"{PARSER_VERSION};pypdf",
        "parse_status": "parsed",
        "machine_readable": True,
        "ocr_required": False,
        "ocr_confidence": None,
        "warnings": [],
        "pages": [],
        "approximate_record_count": 1,
        "date_range": None,
        "currency": None,
        "fields": {},
        "field_labels": [],
    }
    reader = PdfReader(path)
    page_texts: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        warning = None
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # pragma: no cover - source-specific failure
            text = ""
            warning = f"text extraction failed: {type(exc).__name__}: {exc}"
        usable = len(re.sub(r"\s+", "", text)) >= 20
        if not usable:
            profile["machine_readable"] = False
            profile["ocr_required"] = True
            profile["parse_status"] = "partial"
            warning = warning or "No usable text layer; OCR required but not applied automatically."
        profile["pages"].append(
            {
                "page": index,
                "text_characters": len(text),
                "usable_text_layer": usable,
                "warning": warning,
            }
        )
        if warning:
            profile["warnings"].append(f"page {index}: {warning}")
        page_texts.append(text)
    text = "\n".join(page_texts)
    profile["fields"] = extract_inline_fields(text)
    profile["field_labels"] = label_candidates_from_text(text)
    if "RD$" in text or "DOP" in text:
        profile["currency"] = "DOP"
    dates = []
    for match in re.finditer(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b", text):
        try:
            dates.append(date(int(match.group(3)), int(match.group(2)), int(match.group(1))))
        except ValueError:
            continue
    if dates:
        profile["date_range"] = [min(dates).isoformat(), max(dates).isoformat()]
    return profile


def workbook_defined_names(workbook: Any) -> list[dict[str, Any]]:
    names: list[dict[str, Any]] = []
    try:
        iterable = workbook.defined_names.values()
    except Exception:
        iterable = []
    for entry in iterable:
        names.append(
            {
                "name": getattr(entry, "name", None),
                "attr_text": getattr(entry, "attr_text", None),
                "hidden": bool(getattr(entry, "hidden", False)),
            }
        )
    return names


def parse_xlsx(path: Path, relative_path: str) -> dict[str, Any]:
    profile: dict[str, Any] = {
        "relative_path": relative_path,
        "parser": "openpyxl.load_workbook",
        "parser_version": f"{PARSER_VERSION};openpyxl",
        "parse_status": "parsed",
        "machine_readable": True,
        "ocr_required": False,
        "ocr_confidence": None,
        "warnings": [],
        "worksheets": [],
        "approximate_record_count": 1,
        "date_range": None,
        "currency": None,
        "fields": {},
        "field_labels": [],
        "totals": [],
        "line_item_header_rows": [],
        "defined_names": [],
        "external_links": 0,
    }
    wb_formula = load_workbook(path, data_only=False, read_only=False)
    wb_cached = load_workbook(path, data_only=True, read_only=False)
    profile["defined_names"] = workbook_defined_names(wb_formula)
    profile["external_links"] = len(getattr(wb_formula, "_external_links", []) or [])
    all_text: list[str] = []
    all_dates: list[date] = []
    label_candidates: list[str] = []
    extracted_fields: dict[str, list[str]] = defaultdict(list)
    for sheet_index, ws in enumerate(wb_formula.worksheets):
        cached_ws = wb_cached[ws.title]
        sheet_info = {
            "index": sheet_index,
            "title": ws.title,
            "visibility": ws.sheet_state,
            "used_range": ws.calculate_dimension(),
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "merged_ranges": [str(item) for item in ws.merged_cells.ranges],
            "formula_cells": 0,
            "comments": 0,
            "hyperlinks": 0,
            "data_validations": len(getattr(getattr(ws, "data_validations", None), "dataValidation", []) or []),
            "auto_filter": getattr(ws.auto_filter, "ref", None),
            "tables": sorted(ws.tables.keys()),
            "hidden_rows": sum(1 for dimension in ws.row_dimensions.values() if dimension.hidden),
            "hidden_columns": sum(1 for dimension in ws.column_dimensions.values() if dimension.hidden),
            "nonempty_cells": 0,
            "cached_formula_results": 0,
            "cell_samples": [],
        }
        for row in ws.iter_rows():
            row_values: list[tuple[str, Any, Any]] = []
            for cell in row:
                value = cell.value
                cached_value = cached_ws[cell.coordinate].value
                if value is None:
                    continue
                sheet_info["nonempty_cells"] += 1
                if cell.comment is not None:
                    sheet_info["comments"] += 1
                if cell.hyperlink is not None:
                    sheet_info["hyperlinks"] += 1
                if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                    sheet_info["formula_cells"] += 1
                    if cached_value is not None:
                        sheet_info["cached_formula_results"] += 1
                if isinstance(value, (datetime, date)):
                    all_dates.append(value.date() if isinstance(value, datetime) else value)
                if isinstance(value, str):
                    all_text.append(value)
                    label_candidates.extend(label_candidates_from_text(value))
                    if normalize_label(value) in STANDALONE_LABELS:
                        label_candidates.append(re.sub(r"\s+", " ", value).strip())
                    for key, values in extract_inline_fields(value).items():
                        for item in values:
                            if item not in extracted_fields[key]:
                                extracted_fields[key].append(item)
                row_values.append((cell.coordinate, value, cached_value))
            if row_values and len(sheet_info["cell_samples"]) < 40:
                sheet_info["cell_samples"].append(
                    [
                        {
                            "cell": coord,
                            "raw_value": scalar(value),
                            "formula": scalar(value) if isinstance(value, str) and value.startswith("=") else None,
                            "displayed_or_cached_value": scalar(cached_value),
                            "number_format": ws[coord].number_format,
                        }
                        for coord, value, cached_value in row_values[:14]
                    ]
                )
            if row_values:
                rendered = " ".join(str(value) for _, value, _ in row_values if value is not None)
                lowered = normalize_label(rendered)
                if "descripción" in lowered and ("cant" in lowered or "precio" in lowered) and "total" in lowered:
                    profile["line_item_header_rows"].append({"sheet": ws.title, "row": row[0].row, "raw": rendered})
                for coord, raw_value, cached_value in row_values:
                    if not isinstance(raw_value, str):
                        continue
                    normalized = normalize_label(raw_value)
                    for label, destination in TOTAL_LABELS.items():
                        if normalized == label or normalized.startswith(label + " "):
                            next_col = ws.cell(row=ws[coord].row, column=ws[coord].column + 1)
                            next_cached = cached_ws[next_col.coordinate].value
                            amount = parse_decimal(next_cached if next_cached is not None else next_col.value)
                            if amount is not None:
                                profile["totals"].append(
                                    {
                                        "sheet": ws.title,
                                        "label_cell": coord,
                                        "value_cell": next_col.coordinate,
                                        "raw_label": raw_value,
                                        "destination": destination,
                                        "formula": scalar(next_col.value)
                                        if isinstance(next_col.value, str) and next_col.value.startswith("=")
                                        else None,
                                        "displayed_or_cached_value": scalar(next_cached),
                                        "normalized_amount": str(amount),
                                    }
                                )
        profile["worksheets"].append(sheet_info)
    combined_text = "\n".join(all_text)
    for key, values in extract_inline_fields(combined_text).items():
        for item in values:
            if item not in extracted_fields[key]:
                extracted_fields[key].append(item)
    profile["fields"] = dict(extracted_fields)
    profile["field_labels"] = label_candidates
    if "RD$" in combined_text or "DOP" in combined_text:
        profile["currency"] = "DOP"
    if all_dates:
        profile["date_range"] = [min(all_dates).isoformat(), max(all_dates).isoformat()]
    if path.name.upper().startswith("AAAANUMEROS DE FACTURAS"):
        profile["approximate_record_count"] = max(
            (sheet["max_row"] - 1 for sheet in profile["worksheets"]),
            default=0,
        )
    return profile


def parse_one(entry: dict[str, Any]) -> dict[str, Any]:
    path = Path(entry["local_path"])
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            profile = parse_pdf(path, entry["relative_path"])
        elif ext == ".xlsx":
            profile = parse_xlsx(path, entry["relative_path"])
        else:
            return {
                "relative_path": entry["relative_path"],
                "parse_status": entry["parse_status"],
                "parser": None,
                "parser_version": PARSER_VERSION,
                "warnings": [entry.get("reason")] if entry.get("reason") else [],
            }
        return profile
    except Exception as exc:  # pragma: no cover - source-specific failure
        return {
            "relative_path": entry["relative_path"],
            "parse_status": "unreadable",
            "parser": "pypdf.PdfReader" if ext == ".pdf" else "openpyxl.load_workbook",
            "parser_version": PARSER_VERSION,
            "machine_readable": False,
            "ocr_required": ext == ".pdf",
            "ocr_confidence": None,
            "warnings": [f"{type(exc).__name__}: {exc}", traceback.format_exc(limit=3)],
            "fields": {},
            "field_labels": [],
        }


def pair_documents(profiles: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for profile in profiles:
        relative = profile["relative_path"]
        suffix = Path(relative).suffix.lower()
        if suffix not in {".pdf", ".xlsx"}:
            continue
        stem_key = str(Path(relative).with_suffix("")).casefold()
        groups[stem_key][suffix] = profile
    pairs: list[dict[str, Any]] = []
    for stem_key, members in sorted(groups.items()):
        if ".pdf" not in members or ".xlsx" not in members:
            continue
        pdf = members[".pdf"]
        xlsx = members[".xlsx"]
        conflicts = []
        for field in sorted(set(pdf.get("fields", {})) | set(xlsx.get("fields", {}))):
            pdf_values = pdf.get("fields", {}).get(field, [])
            xlsx_values = xlsx.get("fields", {}).get(field, [])
            pdf_norm = {normalize_label(str(item)) for item in pdf_values}
            xlsx_norm = {normalize_label(str(item)) for item in xlsx_values}
            if pdf_norm and xlsx_norm and pdf_norm.isdisjoint(xlsx_norm):
                conflicts.append(
                    {
                        "field": field,
                        "pdf_values": pdf_values,
                        "xlsx_values": xlsx_values,
                        "rule": "normalized value sets are disjoint",
                    }
                )
        pairs.append(
            {
                "stem_key": stem_key,
                "pdf": pdf["relative_path"],
                "xlsx": xlsx["relative_path"],
                "field_conflicts": conflicts,
            }
        )
    return pairs


def build_field_catalog(profiles: list[dict[str, Any]]) -> list[dict[str, Any]]:
    observed: dict[str, dict[str, Any]] = {}
    for profile in profiles:
        relative = profile["relative_path"]
        for raw_label in profile.get("field_labels", []):
            normalized = normalize_label(raw_label)
            if not normalized:
                continue
            entry = observed.setdefault(
                normalized,
                {
                    "normalized_source_field": normalized,
                    "source_labels": Counter(),
                    "source_locations": [],
                    "examples": [],
                    "source_formats": Counter(),
                    "occurrences": 0,
                },
            )
            entry["source_labels"][raw_label] += 1
            entry["occurrences"] += 1
            if len(entry["source_locations"]) < 12:
                entry["source_locations"].append(relative)
            if len(entry["examples"]) < 8 and raw_label not in entry["examples"]:
                entry["examples"].append(raw_label)
            entry["source_formats"][Path(relative).suffix.lower()] += 1
        for field, values in profile.get("fields", {}).items():
            normalized = f"parsed::{field}"
            entry = observed.setdefault(
                normalized,
                {
                    "normalized_source_field": normalized,
                    "source_labels": Counter(),
                    "source_locations": [],
                    "examples": [],
                    "source_formats": Counter(),
                    "occurrences": 0,
                },
            )
            entry["source_labels"][field] += len(values) or 1
            entry["occurrences"] += len(values) or 1
            if len(entry["source_locations"]) < 12:
                entry["source_locations"].append(relative)
            for value in values:
                if len(entry["examples"]) < 8 and value not in entry["examples"]:
                    entry["examples"].append(value)
            entry["source_formats"][Path(relative).suffix.lower()] += 1
    result = []
    for key, entry in sorted(observed.items()):
        result.append(
            {
                "normalized_source_field": key,
                "source_labels": dict(entry["source_labels"].most_common()),
                "source_locations": entry["source_locations"],
                "examples": entry["examples"],
                "source_formats": dict(entry["source_formats"]),
                "occurrences": entry["occurrences"],
            }
        )
    return result


def manifest_summary(manifest: list[dict[str, Any]], profiles: list[dict[str, Any]]) -> dict[str, Any]:
    by_relative = {profile["relative_path"]: profile for profile in profiles}
    counts = Counter()
    extensions = Counter()
    for entry in manifest:
        profile = by_relative.get(entry["relative_path"], {})
        status = profile.get("parse_status", entry["parse_status"])
        counts[status] += 1
        extensions[entry["extension"] or "[none]"] += 1
    reconciliation = {
        "total": len(manifest),
        "parsed": counts["parsed"],
        "partial": counts["partial"],
        "unreadable": counts["unreadable"],
        "unsupported": counts["unsupported"],
        "irrelevant": counts["restricted_inventory_only"] + counts["irrelevant"],
    }
    reconciliation["equation_holds"] = reconciliation["total"] == sum(
        reconciliation[key] for key in ("parsed", "partial", "unreadable", "unsupported", "irrelevant")
    )
    hashes: dict[str, list[str]] = defaultdict(list)
    for entry in manifest:
        hashes[entry["sha256"]].append(entry["relative_path"])
    duplicate_groups = [paths for paths in hashes.values() if len(paths) > 1]
    return {
        "parser_version": PARSER_VERSION,
        "generated_at": datetime.now().astimezone().isoformat(),
        "reconciliation": reconciliation,
        "extensions": dict(extensions),
        "duplicate_hash_groups": duplicate_groups,
        "exact_duplicate_file_count": sum(len(group) - 1 for group in duplicate_groups),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--acquired-root", type=Path, required=True)
    parser.add_argument("--matrix", type=Path, required=True)
    parser.add_argument("--restricted-root", type=Path, required=True)
    parser.add_argument("--archive-root", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    parser.add_argument("--workers", type=int, default=6)
    args = parser.parse_args()

    discovered = discover_source_paths(args.acquired_root, args.matrix, args.restricted_root)
    manifest = [base_manifest(path, relative) for path, relative in discovered]
    write_json(args.output_root / "manifest-initial.json", manifest)

    archives = []
    for archive in sorted(args.archive_root.glob("*")):
        if archive.is_file():
            stat = archive.stat()
            archives.append(
                {
                    "file_name": archive.name,
                    "size_bytes": stat.st_size,
                    "downloaded_at": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(),
                    "sha256": sha256_file(archive),
                }
            )
    write_json(args.output_root / "acquisition-archives.json", archives)

    parse_entries = [entry for entry in manifest if entry["extension"] in INVOICE_EXTENSIONS]
    profiles: list[dict[str, Any]] = []
    with concurrent.futures.ThreadPoolExecutor(max_workers=max(1, args.workers)) as executor:
        futures = {executor.submit(parse_one, entry): entry for entry in parse_entries}
        for index, future in enumerate(concurrent.futures.as_completed(futures), start=1):
            profiles.append(future.result())
            if index % 100 == 0:
                print(f"parsed {index}/{len(parse_entries)}", flush=True)
    for entry in manifest:
        if entry["extension"] not in INVOICE_EXTENSIONS:
            profiles.append(parse_one(entry))
    profiles.sort(key=lambda item: item["relative_path"].casefold())

    profile_by_path = {profile["relative_path"]: profile for profile in profiles}
    for entry in manifest:
        profile = profile_by_path.get(entry["relative_path"])
        if profile:
            entry["parse_status"] = profile.get("parse_status", entry["parse_status"])
            if profile.get("warnings"):
                entry["reason"] = "; ".join(str(item) for item in profile["warnings"][:3])

    pairs = pair_documents(profiles)
    field_catalog = build_field_catalog(profiles)
    summary = manifest_summary(manifest, profiles)
    pdf_text_counts = [
        sum(page.get("text_characters", 0) for page in profile.get("pages", []))
        for profile in profiles
        if Path(profile["relative_path"]).suffix.lower() == ".pdf"
    ]
    summary["paired_pdf_xlsx_documents"] = len(pairs)
    summary["pair_field_conflict_count"] = sum(len(pair["field_conflicts"]) for pair in pairs)
    summary["pdf_text_character_median"] = statistics.median(pdf_text_counts) if pdf_text_counts else 0
    summary["field_catalog_entries"] = len(field_catalog)

    write_json(args.output_root / "manifest-initial.json", manifest)
    write_json(args.output_root / "file-profiles.json", profiles)
    write_json(args.output_root / "document-pairs.json", pairs)
    write_json(args.output_root / "field-catalog.json", field_catalog)
    write_json(args.output_root / "summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
